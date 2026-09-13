"""CE QUE FAIT LE CODE :
Exporte la source narrative en Word sans modifier la source ni les archives.
Prérequis de cet export facultatif : Python 3 et python-docx.
Les vues Markdown, elles, ne demandent que Node.js.
"""
from __future__ import annotations
import hashlib
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError as exc:
    raise SystemExit("Export Word facultatif : installer python-docx avec python -m pip install python-docx (ou py -3 -m pip install python-docx sur Windows).") from exc

ROOT=Path(__file__).resolve().parent.parent
source=ROOT/'STORY_SOURCE.md'
raw=source.read_text(encoding='utf-8').replace('\r\n','\n')
digest=hashlib.sha256(raw.encode('utf-8')).hexdigest()
vm=re.search(r'^Version : (.*?) — (.*?)$',raw,re.M)
if not vm: raise SystemExit('Version et date manquantes dans STORY_SOURCE.md.')
version,date=vm.groups()
pattern=re.compile(r'<!-- BEGIN:([A-Z0-9-]+) -->\n<!-- META:(.+?) -->\n(?:<a id="[^"]+"></a>\n)?## ([^\n]+)\n([\s\S]*?)<!-- END:\1 -->')
blocks=[]
for m in pattern.finditer(raw):
    id,meta,title,text=m.groups()
    group=list(re.finditer(r'^# (.+)$',raw[:m.start()],re.M))[-1].group(1)
    blocks.append(dict(id=id,meta=json.loads(meta),title=title,text=text.strip(),group=group))
ids=[b['id'] for b in blocks]
if not blocks or len(set(ids))!=len(ids): raise SystemExit('Blocs absents ou identifiants dupliqués.')
if len(re.findall(r'^<!-- BEGIN:',raw,re.M))!=len(blocks):raise SystemExit('Bloc mal formé.')
for b in blocks:
    if any(x.strip() not in ids for x in b['meta']['refs'].split(',') if x.strip()):raise SystemExit('Référence inconnue dans '+b['id'])

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
sec.left_margin=Inches(.77);sec.right_margin=Inches(.77)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.32)
usable=sec.page_width-sec.left_margin-sec.right_margin
styles=doc.styles
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[n].font.name='Calibri'
normal=styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.widow_control=True
for n,size,color in [('Title',34,'163749'),('Heading 1',23,'163749'),('Heading 2',14,'163749'),('Heading 3',11.5,'22696B')]:
    st=styles[n];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
    st.font.bold=True;st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(7)
styles['Subtitle'].font.size=Pt(16)
styles['Caption'].font.size=Pt(8.5)
styles['Caption'].font.color.rgb=RGBColor.from_string('5B6870')
styles['Caption'].paragraph_format.space_after=Pt(5)

header=sec.header.paragraphs[0]
header.text='LES ORIGINES  /  RÉFÉRENCE NARRATIVE'
header.style=styles['Caption']
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run(f'Export v{version} • Source : STORY_SOURCE.md  |  ')
run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('5B6870')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title=f'Les Origines — Bible narrative v{version}'
doc.core_properties.subject='Référence narrative générée depuis une source Markdown unique'
doc.core_properties.author='Projet de jeu original'
doc.core_properties.keywords='Orthe, Elio, Lyra, Nacre, narration, propositions, chronologie'

p=doc.add_paragraph('PROJET DE JEU ORIGINAL',style='Caption')
p.paragraph_format.space_before=Pt(70)
doc.add_paragraph('LES ORIGINES',style='Title')
doc.add_paragraph('Bible narrative et registre de décisions',style='Subtitle')
doc.add_paragraph(f'Version {version} • {date}')
p=doc.add_paragraph('Elio et Lyra\nLe noyau intact, le Concordat et la dernière fenêtre')
p.paragraph_format.space_before=Pt(28);p.paragraph_format.space_after=Pt(22)
for r in p.runs:r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string('22696B')
doc.add_paragraph('Refonte de la version 1.0 intégrant les précisions de l’auteur. Les décisions confirmées, les propositions et les questions restent distinguées. Divulgâchage intégral.')
doc.add_paragraph('Fichier de travail : STORY_SOURCE.md. Ce document Word est un instantané de lecture ; il ne doit pas devenir une seconde source à maintenir.',style='Normal')
p=doc.add_paragraph('Empreinte de la source : '+digest[:16]+'…',style='Caption')
p.paragraph_format.space_before=Pt(30)
doc.add_page_break()
doc.add_heading('Repères de lecture',level=1)
doc.add_paragraph('Le scénario de ce dossier est écrit avec Elio comme protagoniste. Le choix de Lyra inverse les rôles dans toute l’histoire, sans créer une seconde intrigue.')
groups=list(dict.fromkeys(b['group'] for b in blocks))
for i,g in enumerate(groups,1):doc.add_paragraph(f'{i:02d}   {g}')
doc.add_heading('Statuts et mise à jour',level=2)
doc.add_paragraph('CONFIRME : orientation fixée par l’auteur. PROPOSE : solution de travail issue de l’auteur, de l’assistant ou de la v1.0. OUVERT : arbitrage non résolu. Une précision proposée à l’intérieur d’un bloc n’est pas confirmée par simple voisinage.')
doc.add_paragraph('Chaque bloc possède un identifiant stable. Pour corriger le dossier, modifier ce bloc dans STORY_SOURCE.md, puis régénérer les vues. Les références à la fin des blocs indiquent les dépendances à relire.')
doc.add_paragraph('Une règle temporelle modifiée peut demander de revoir plusieurs scènes. Le générateur synchronise les documents et vérifie les liens ; il ne remplace pas le travail de cohérence narrative.')

bookmark_no=0

def inline(p,text):
    parts=re.split(r'(\*\*.*?\*\*|`[^`]*`)',text)
    for part in parts:
        if not part:continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]);r.bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        else:p.add_run(part)

def paragraph(text,style=None):
    p=doc.add_paragraph(style=style)
    inline(p,text)
    return p

def table(lines):
    rows=[]
    for line in lines:
        parts=[x.strip() for x in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-+:?',x.replace(' ','')) for x in parts):continue
        rows.append(parts)
    if not rows:return
    cols=len(rows[0]); t=doc.add_table(rows=1,cols=cols);t.autofit=False
    if cols==2:fractions=[.29,.71]
    elif cols==3:fractions=[.23,.35,.42]
    else:fractions=[1/cols]*cols
    for col,fraction in zip(t.columns,fractions):col.width=int(usable*fraction)
    for ri,values in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,(cell,text) in enumerate(zip(cells,values)):
            cell.width=int(usable*fractions[ci])
            cell.vertical_alignment=1
            tcPr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',80),('bottom',80),('left',90),('right',90)]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');margins.append(x)
            tcPr.append(margins)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'163749' if ri==0 else ('F0F5F5' if ri%2 else 'FAFBFC'));tcPr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.04
            inline(p,text)
            for r in p.runs:
                r.font.size=Pt(9.3)
                if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        trPr=t.rows[ri]._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit');trPr.append(ns)
        if ri==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def body(text):
    lines=text.splitlines();i=0;buf=[]
    def flush():
        if buf:paragraph(' '.join(buf));buf.clear()
    while i<len(lines):
        line=lines[i].strip()
        if not line:flush();i+=1;continue
        if line.startswith('|'):
            flush();ts=[]
            while i<len(lines) and lines[i].strip().startswith('|'):ts.append(lines[i]);i+=1
            table(ts);continue
        if line.startswith('### '):flush();doc.add_heading(line[4:],level=3);i+=1;continue
        buf.append(line);i+=1
    flush()

last=None
for b in blocks:
    if b['group']!=last:
        
        if last is None: doc.add_page_break()
        doc.add_heading(b['group'],level=1);last=b['group']
    label=b['title']
    if b['meta']['kind']=='chapitre':label=b['meta']['period']+' — '+label
    hp=doc.add_heading(label,level=2)
    bookmark_no+=1
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_no));start.set(qn('w:name'),b['id'].replace('-','_'))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_no))
    hp._p.insert(0,start);hp._p.append(end)
    stat=f"{b['id']}  ·  {b['meta']['status']}  ·  origine : {b['meta']['origin']}"
    if b['meta']['kind']=='evenement':stat+='  ·  '+b['meta']['period']
    p=doc.add_paragraph(stat,style='Caption');p.paragraph_format.keep_with_next=True
    body(b['text'])
    refs=[r.strip() for r in b['meta']['refs'].split(',') if r.strip()]
    if refs:
        p=doc.add_paragraph('Blocs associés : ',style='Caption')
        for i,ref in enumerate(refs):
            if i:p.add_run('  ·  ')
            link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),ref.replace('-','_'))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            col=OxmlElement('w:color');col.set(qn('w:val'),'22696B');pr.append(col)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');pr.append(sz);r.append(pr)
            tx=OxmlElement('w:t');tx.text=ref;r.append(tx);link.append(r);p._p.append(link)

out=ROOT/'exports';out.mkdir(parents=True,exist_ok=True)
filename=f"Jeu_original_Bible_narrative_v{version.replace('.','_')}_Monde_des_Origines.docx"
path=out/filename
doc.save(path)
(out/'EXPORT.json').write_text(json.dumps({'version':version,'date':date,'file':filename,'source_sha256':digest,'source':'STORY_SOURCE.md'},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(path)
